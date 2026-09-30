"""Bounded PDF/DOCX text extraction."""

import re
import zipfile
from pathlib import Path

import pymupdf as fitz
from docx import Document

from config import (
    MAX_DOCX_COMPRESSION_RATIO,
    MAX_DOCX_UNCOMPRESSED_BYTES,
    MAX_EXTRACTED_TEXT_CHARS,
)


class ResumeParseError(ValueError):
    """Raised when an uploaded resume cannot be read safely."""


def clean_text(text: str) -> str:
    text = text.replace("\x00", " ").replace("\r", "\n")
    lines = [re.sub(r"[ \t]+", " ", line).strip() for line in text.splitlines()]
    return "\n".join(line for line in lines if line).strip()[:MAX_EXTRACTED_TEXT_CHARS]


def extract_pdf_text(path) -> str:
    try:
        pdf_bytes = Path(path).read_bytes()
        if pdf_bytes[:5] != b"%PDF-":
            raise ResumeParseError("This file is not a valid PDF.")
        with fitz.open(stream=pdf_bytes, filetype="pdf") as document:
            if document.needs_pass:
                raise ResumeParseError("Password-protected PDFs are not supported.")
            text_parts, size = [], 0
            for page in document:
                content = page.get_text("text")
                text_parts.append(content[: max(0, MAX_EXTRACTED_TEXT_CHARS - size)])
                size += len(content)
                if size >= MAX_EXTRACTED_TEXT_CHARS:
                    break
    except ResumeParseError:
        raise
    except Exception as exc:
        raise ResumeParseError("This PDF is corrupted or could not be read.") from exc
    cleaned = clean_text("\n".join(text_parts))
    if not cleaned:
        raise ResumeParseError("Unable to extract readable text from this PDF. Please upload a text-based resume PDF.")
    return cleaned


def _validate_docx_archive(path) -> None:
    try:
        with zipfile.ZipFile(path) as archive:
            entries = archive.infolist()
            total = sum(entry.file_size for entry in entries)
            compressed = sum(max(entry.compress_size, 1) for entry in entries)
            if total > MAX_DOCX_UNCOMPRESSED_BYTES or total / max(compressed, 1) > MAX_DOCX_COMPRESSION_RATIO:
                raise ResumeParseError("This DOCX archive expands beyond safe processing limits.")
            if "word/document.xml" not in {entry.filename for entry in entries}:
                raise ResumeParseError("This file is not a valid DOCX document.")
    except ResumeParseError:
        raise
    except (zipfile.BadZipFile, OSError) as exc:
        raise ResumeParseError("This DOCX file is corrupted or could not be read.") from exc


def extract_docx_text(path) -> str:
    _validate_docx_archive(path)
    try:
        document = Document(path)
        parts = [paragraph.text for paragraph in document.paragraphs]
        parts.extend(cell.text for table in document.tables for row in table.rows for cell in row.cells)
    except Exception as exc:
        raise ResumeParseError("This DOCX file is corrupted or could not be read.") from exc
    cleaned = clean_text("\n".join(parts))
    if not cleaned:
        raise ResumeParseError("Unable to extract readable text from this DOCX file.")
    return cleaned


def extract_resume_text(path) -> str:
    extension = Path(path).suffix.lower()
    if extension == ".pdf":
        return extract_pdf_text(path)
    if extension == ".docx":
        return extract_docx_text(path)
    raise ResumeParseError("Only PDF and DOCX resumes are supported.")
