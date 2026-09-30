import io
import sqlite3
import zipfile
from pathlib import Path

import pymupdf as fitz
import pytest
from docx import Document

from app import create_app
from services.resume_parser import ResumeParseError, extract_docx_text
from services.resume_scorer import has_valid_phone, measurable_outcomes, score_resume
from services.skill_extractor import extract_skills, flatten_skills


def skills(text):
    return set(flatten_skills(extract_skills(text)))


def make_pdf(text):
    document = fitz.open()
    page = document.new_page()
    page.insert_text((72, 72), text)
    data = document.tobytes()
    document.close()
    return data


def test_aliases_versions_and_technology_boundaries():
    found = skills("Technical Skills: HTML5 CSS3 Python3 Java8 OOPs AngularJS k8s sklearn")
    assert {"HTML", "CSS", "Python", "Java", "OOP", "AngularJS", "Kubernetes", "Scikit-learn"} <= found
    assert "Java" not in skills("JavaScript React.js Node.js")
    assert {"C", "C++", "C#"} <= skills("Languages: C, C++, C#")
    assert skills("MySQL PostgreSQL") == {"MySQL", "PostgreSQL"}


def test_ambiguous_false_positives_are_rejected():
    assert "C" not in skills("Vitamin C supports nutrition")
    assert "R" not in skills("Worked in R&D")
    assert "Spring" not in skills("Graduating Spring 2024\nTechnical Skills\nPython")
    assert "Go" not in skills("I go to the gym")


def test_heading_variants_need_substantive_content():
    for heading in ("Professional Experience", "Work Experience", "Internship Experience", "Technical Skills", "Key Skills"):
        result = score_resume(f"Name\n{heading}\n\nEducation\n", {})
        assert result["total_score"] < 20


def test_phone_and_metrics_regex_regressions():
    assert not has_valid_phone("Education 2019 - 2023")
    assert has_valid_phone("+91 98765 43210")
    values = measurable_outcomes("Improved 20%, then +35%, reached 40% and 2.5x throughput")
    assert len(values) == 4


def test_docx_tables_are_extracted(tmp_path):
    path = tmp_path / "table.docx"
    doc = Document()
    table = doc.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Technical Skills"
    table.cell(0, 1).text = "Python Docker"
    doc.save(path)
    assert "Python Docker" in extract_docx_text(path)


def test_docx_high_compression_ratio_is_rejected(tmp_path):
    path = tmp_path / "bomb.docx"
    with zipfile.ZipFile(path, "w", compression=zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("word/document.xml", "0" * 1_000_000)
    with pytest.raises(ResumeParseError, match="safe processing limits"):
        extract_docx_text(path)


def test_api_validation_ownership_multiple_matches_and_delete(app, client):
    assert client.post("/api/match-job", data="{bad", content_type="application/json").status_code == 400
    assert client.post("/api/match-job", json=[]).status_code == 400
    assert client.post("/api/match-job", json={"analysis_id": {}, "job_description": []}).status_code == 400
    response = client.post("/api/analyze-resume", data={"resume": (io.BytesIO(make_pdf("x@example.com Python")), "履歴書.pdf", "application/pdf")}, content_type="multipart/form-data")
    assert response.status_code == 201
    analysis_id = response.get_json()["id"]
    for description in ("Python Docker", "Python AWS"):
        assert client.post("/api/match-job", json={"analysis_id": analysis_id, "job_description": description}).status_code == 201
    db = sqlite3.connect(app.config["DATABASE"])
    try:
        assert db.execute("SELECT COUNT(*) FROM job_matches").fetchone()[0] == 2
    finally:
        db.close()
    other = app.test_client()
    assert other.get(f"/api/history/{analysis_id}").status_code == 404
    assert client.delete(f"/api/history/{analysis_id}").status_code == 200
    assert client.get(f"/api/history/{analysis_id}").status_code == 404


def test_job_description_limits_and_origin_protection(client):
    demo = client.post("/api/demo").get_json()
    too_long = "x" * 20001
    assert client.post("/api/match-job", json={"analysis_id": demo["id"], "job_description": too_long}).status_code == 422
    assert client.post("/api/demo", headers={"Origin": "https://evil.example"}).status_code == 403


def test_database_has_history_and_foreign_key_indexes(app):
    db = sqlite3.connect(app.config["DATABASE"])
    try:
        analysis_indexes = {row[1] for row in db.execute("PRAGMA index_list(analyses)")}
        match_indexes = {row[1] for row in db.execute("PRAGMA index_list(job_matches)")}
        foreign_keys = db.execute("PRAGMA foreign_key_list(job_matches)").fetchall()
    finally:
        db.close()
    assert "idx_analyses_owner_created" in analysis_indexes
    assert "idx_job_matches_analysis" in match_indexes
    assert any(row[2] == "analyses" and row[6].upper() == "CASCADE" for row in foreign_keys)


def test_secret_key_fails_closed_except_for_explicit_local_or_test_mode(monkeypatch, tmp_path):
    monkeypatch.delenv("SECRET_KEY", raising=False)
    monkeypatch.delenv("RESUMEIQ_ENV", raising=False)
    with pytest.raises(RuntimeError, match="SECRET_KEY"):
        create_app({"DATABASE": str(tmp_path / "unset.db"), "UPLOAD_FOLDER": str(tmp_path / "unset")})
    monkeypatch.setenv("RESUMEIQ_ENV", "production")
    with pytest.raises(RuntimeError, match="SECRET_KEY"):
        create_app({"DATABASE": str(tmp_path / "prod.db"), "UPLOAD_FOLDER": str(tmp_path / "prod")})
    monkeypatch.setenv("RESUMEIQ_ENV", "development")
    local = create_app({"DATABASE": str(tmp_path / "local.db"), "UPLOAD_FOLDER": str(tmp_path / "local")})
    assert "development" in local.config["SECRET_KEY"]


def test_ci_runs_javascript_tests():
    workflow = (Path(__file__).resolve().parents[1] / ".github" / "workflows" / "ci.yml").read_text(encoding="utf-8")
    assert "node --test static/js/app_helpers.test.js" in workflow
